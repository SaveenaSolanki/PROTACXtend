#!/usr/bin/env python3
"""Build the ProtacXtend experimental-study deck (figure-forward, deep).

Embeds the real figures produced by study/make_figures.py and carries the full
technical content of the Phase 0 package (00-03).

Run:
    python study/make_figures.py        # regenerate figures first
    python study/build_presentation.py  # then the deck
Output:
    study/PROTACXTEND_STUDY_DECK.pptx
"""
from __future__ import annotations

from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.util import Emu, Inches, Pt

ROOT = Path(__file__).resolve().parent
FIG = ROOT / "figures"
OUT = ROOT / "PROTACXTEND_STUDY_DECK.pptx"

# ── palette ───────────────────────────────────────────────────────────────
NAVY = RGBColor(0x0F, 0x1E, 0x33)
INK = RGBColor(0x1B, 0x24, 0x30)
SLATE = RGBColor(0x4A, 0x57, 0x68)
TEAL = RGBColor(0x17, 0xA3, 0x98)
TEAL_L = RGBColor(0xE3, 0xF3, 0xF0)
AMBER = RGBColor(0xE8, 0xA3, 0x3D)
AMBER_L = RGBColor(0xFB, 0xF0, 0xDD)
RED = RGBColor(0xC4, 0x45, 0x3B)
RED_L = RGBColor(0xFB, 0xE9, 0xE7)
PAPER = RGBColor(0xF4, 0xF6, 0xF9)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LINE = RGBColor(0xD5, 0xDB, 0xE3)
SOFT = RGBColor(0xE9, 0xEE, 0xF4)
MUTE = RGBColor(0xBE, 0xCB, 0xD6)

FONT = "Arial"
W, H = Inches(13.333), Inches(7.5)

prs = Presentation()
prs.slide_width = W
prs.slide_height = H
BLANK = prs.slide_layouts[6]
_n = {"i": 0}


# ── primitives ────────────────────────────────────────────────────────────
def _tb(slide, x, y, w, h, text, *, size=16, color=INK, bold=False,
        align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, italic=False, spacing=1.0):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = spacing
    r = p.add_run()
    r.text = text
    r.font.name = FONT
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = color
    return box


def _rect(slide, x, y, w, h, fill, *, line=None, radius=None, shape=MSO_SHAPE.ROUNDED_RECTANGLE):
    s = slide.shapes.add_shape(shape, x, y, w, h)
    s.fill.solid()
    s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(1)
    s.shadow.inherit = False
    if radius is not None and shape == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    return s


def _footer(slide, section, dark=False):
    _n["i"] += 1
    bg = RGBColor(0x0A, 0x15, 0x24) if dark else PAPER
    fg = MUTE if dark else SLATE
    _rect(slide, 0, H - Inches(0.32), W, Inches(0.32), bg, shape=MSO_SHAPE.RECTANGLE)
    _tb(slide, Inches(0.45), H - Inches(0.29), Inches(10.5), Inches(0.26),
        f"ProtacXtend Experimental Study   ·   PROTACXTEND-STUDY/1.0.0   ·   {section}",
        size=8.5, color=fg)
    _tb(slide, W - Inches(1.0), H - Inches(0.29), Inches(0.6), Inches(0.26),
        f"{_n['i']:02d}", size=8.5, color=fg, align=PP_ALIGN.RIGHT, bold=True)


def content_slide(title, kicker=""):
    slide = prs.slides.add_slide(BLANK)
    _rect(slide, 0, 0, W, Inches(0.09), TEAL, shape=MSO_SHAPE.RECTANGLE)
    if kicker:
        _tb(slide, Inches(0.5), Inches(0.24), Inches(12.3), Inches(0.3),
            kicker.upper(), size=10.5, color=TEAL, bold=True)
        _tb(slide, Inches(0.5), Inches(0.5), Inches(12.3), Inches(0.66),
            title, size=25, color=NAVY, bold=True)
    else:
        _tb(slide, Inches(0.5), Inches(0.38), Inches(12.3), Inches(0.75),
            title, size=26, color=NAVY, bold=True)
    _footer(slide, kicker or "")
    return slide


def bullets(slide, items, *, x=Inches(0.6), y=Inches(1.5), w=Inches(12.1),
            h=Inches(5.3), size=15, gap=9):
    box = slide.shapes.add_textbox(x, y, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for it in items:
        text, lvl = it if isinstance(it, tuple) else (it, 0)
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.space_after = Pt(gap)
        p.level = lvl
        r = p.add_run()
        r.text = ("•  " if lvl == 0 else "–  ") + text
        r.font.name = FONT
        r.font.size = Pt(size if lvl == 0 else size - 2)
        r.font.color.rgb = INK if lvl == 0 else SLATE
    return box


def table(slide, headers, rows, *, x=Inches(0.5), y=Inches(1.45), w=Inches(12.33),
          h=Inches(4.9), col_widths=None, fsize=11, header_fill=NAVY, row_h=None,
          align=None):
    n_rows, n_cols = len(rows) + 1, len(headers)
    shp = slide.shapes.add_table(n_rows, n_cols, x, y, w, h)
    tbl = shp.table
    if col_widths:
        tot = sum(col_widths)
        for i, cw in enumerate(col_widths):
            tbl.columns[i].width = Emu(int(w * cw / tot))
    if row_h:
        for i in range(n_rows):
            tbl.rows[i].height = row_h
    for j, htxt in enumerate(headers):
        c = tbl.cell(0, j)
        c.text = htxt
        c.fill.solid()
        c.fill.fore_color.rgb = header_fill
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        c.margin_left = Inches(0.07); c.margin_right = Inches(0.05)
        for p in c.text_frame.paragraphs:
            for r in p.runs:
                r.font.name = FONT; r.font.size = Pt(fsize); r.font.bold = True
                r.font.color.rgb = WHITE
    for i, row in enumerate(rows, start=1):
        for j, val in enumerate(row):
            c = tbl.cell(i, j)
            c.text = str(val)
            c.fill.solid()
            c.fill.fore_color.rgb = SOFT if i % 2 == 0 else WHITE
            c.vertical_anchor = MSO_ANCHOR.MIDDLE
            c.margin_left = Inches(0.07); c.margin_right = Inches(0.05)
            for p in c.text_frame.paragraphs:
                if align:
                    p.alignment = align[j]
                for r in p.runs:
                    r.font.name = FONT; r.font.size = Pt(fsize); r.font.color.rgb = INK
    return tbl


def _place_image(slide, name, ax, ay, aw, ah):
    """Fit image inside the box (ax,ay,aw,ah), centered."""
    path = FIG / f"{name}.png"
    iw, ih = Image.open(path).size
    ar = iw / ih
    tw, th = aw, int(aw / ar)
    if th > ah:
        th, tw = ah, int(ah * ar)
    x = ax + int((aw - tw) / 2)
    y = ay + int((ah - th) / 2)
    slide.shapes.add_picture(str(path), x, y, width=tw, height=th)
    return x, y, tw, th


def figure_slide(title, kicker, image, caption, bullets_text=None, *,
                 layout="full", side="left", source="", source_color=SLATE,
                 cap_size=12, bul_size=13.5, note=""):
    slide = content_slide(title, kicker)
    if layout == "full":
        ax, ay, aw, ah = Inches(0.5), Inches(1.3), Inches(12.33), Inches(4.35)
        _place_image(slide, image, ax, ay, aw, ah)
        if source:
            _tb(slide, Inches(0.55), Inches(5.7), Inches(12.2), Inches(0.3),
                source, size=9.5, color=source_color, bold=True)
        _tb(slide, Inches(0.55), Inches(6.0), Inches(12.2), Inches(0.9),
            caption, size=cap_size, color=INK, spacing=1.15)
        if bullets_text:
            _tb(slide, Inches(0.55), Inches(6.7), Inches(12.2), Inches(0.7),
                bullets_text, size=11.5, color=SLATE, spacing=1.1)
    else:
        if side == "left":
            ax, ay, aw, ah = Inches(0.5), Inches(1.4), Inches(5.95), Inches(4.85)
            tx, tw = Inches(6.75), Inches(6.1)
        else:
            ax, ay, aw, ah = Inches(6.9), Inches(1.4), Inches(5.95), Inches(4.85)
            tx, tw = Inches(0.55), Inches(6.1)
        _place_image(slide, image, ax, ay, aw, ah)
        if source:
            _tb(slide, ax, Inches(6.3), aw, Inches(0.3), source, size=9, color=source_color,
                bold=True, align=PP_ALIGN.CENTER)
        if bullets_text:
            bullets(slide, bullets_text, x=tx, y=Inches(1.55), w=tw, h=Inches(4.9),
                    size=bul_size, gap=10)
        _tb(slide, tx, Inches(6.55), tw, Inches(0.6), caption, size=10.5,
            color=SLATE, italic=True, spacing=1.1)
    notes(slide, note or f"{title}. {caption} {bullets_text or ''}".strip())
    return slide


def notes(slide, text):
    slide.notes_slide.notes_text_frame.text = text


def section_slide(num, title, subtitle, note=""):
    slide = prs.slides.add_slide(BLANK)
    _rect(slide, 0, 0, W, H, NAVY, shape=MSO_SHAPE.RECTANGLE)
    _rect(slide, 0, 0, Inches(0.2), H, TEAL, shape=MSO_SHAPE.RECTANGLE)
    _tb(slide, Inches(0.95), Inches(2.5), Inches(11.5), Inches(1.2), num,
        size=64, color=TEAL, bold=True)
    _tb(slide, Inches(0.95), Inches(3.7), Inches(11.5), Inches(1.0), title,
        size=34, color=WHITE, bold=True)
    _tb(slide, Inches(0.95), Inches(4.7), Inches(11.5), Inches(0.8), subtitle,
        size=16, color=MUTE)
    _footer(slide, "section", dark=True)
    notes(slide, note or subtitle)
    return slide


# ══════════════════════════════════════════════════════════════════════════
# 01 TITLE
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
_rect(s, 0, 0, W, H, NAVY, shape=MSO_SHAPE.RECTANGLE)
_rect(s, 0, 0, Inches(0.22), H, TEAL, shape=MSO_SHAPE.RECTANGLE)
# decorative strip of embedded mini-figures
try:
    _place_image(s, "f08_closed48_outcomes", Inches(7.65), Inches(1.05), Inches(5.2), Inches(3.0))
    _place_image(s, "f06_architecture", Inches(7.65), Inches(4.2), Inches(5.2), Inches(1.9))
except Exception:
    pass
_tb(s, Inches(0.9), Inches(1.5), Inches(6.6), Inches(1.4),
    "ProtacXtend\nExperimental Study", size=40, color=WHITE, bold=True, spacing=1.05)
_tb(s, Inches(0.92), Inches(3.35), Inches(6.4), Inches(1.4),
    "A six-hypothesis, evidence-grounded evaluation of a TPD-specific scientific agent",
    size=17, color=RGBColor(0xBF, 0xD4, 0xDA), spacing=1.2)
_rect(s, Inches(0.92), Inches(4.75), Inches(3.0), Inches(0.05), AMBER, shape=MSO_SHAPE.RECTANGLE)
_tb(s, Inches(0.92), Inches(5.0), Inches(6.4), Inches(1.2),
    "Phase 0 — repository freeze · system audit · benchmark audit · pre-execution design\n"
    "Study id PROTACXTEND-STUDY/1.0.0  ·  frozen at 82a0e4d (sprint-2)\n"
    "Status: design complete — pilot awaiting approval",
    size=12, color=MUTE, spacing=1.25)
notes(s, "Title. Right side shows real figures from the audit: closed48 outcomes and the architecture.")
_n["i"] += 1

# ══════════════════════════════════════════════════════════════════════════
# 02 EXECUTIVE SUMMARY
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("What we found and what we will run", "Executive summary")
cards = [
    ("48", "governed benchmark cases", "19 objectively scorable · 29 need expert review · 0 adjudicated", RED),
    ("1,000", "question-bank tasks", "benchmark500 (general + temporal), 0 curated gold", RED),
    ("27", "TPD capabilities", "10 READY · 6 EXECUTABLE · 11 BLOCKED", AMBER),
    ("144", "real runs already on disk", "48 cases × 3 arms, correctness pending gold", TEAL),
    ("17,040", "runs planned", "4 systems × populations × replicates", NAVY),
    ("~$147", "est. API budget", "≈1,530 GPU-h dominates the true cost", NAVY),
]
x0, y0 = Inches(0.5), Inches(1.5)
for i, (big, mid, sub, col) in enumerate(cards):
    c, r = i % 3, i // 3
    x = x0 + c * Inches(4.14)
    y = y0 + r * Inches(2.15)
    _rect(s, x, y, Inches(3.94), Inches(1.95), PAPER, line=LINE, radius=0.06)
    _rect(s, x, y, Inches(0.14), Inches(1.95), col, radius=0.4)
    _tb(s, x + Inches(0.3), y + Inches(0.18), Inches(3.5), Inches(0.7), big,
        size=30, color=col, bold=True)
    _tb(s, x + Inches(0.3), y + Inches(0.95), Inches(3.5), Inches(0.4), mid,
        size=12.5, color=NAVY, bold=True)
    _tb(s, x + Inches(0.3), y + Inches(1.35), Inches(3.5), Inches(0.5), sub,
        size=10.5, color=SLATE, spacing=1.1)
_tb(s, Inches(0.5), Inches(5.95), Inches(12.3), Inches(0.9),
    "Blocker: correctness cannot be scored until gold is adjudicated (primary endpoint denominator = 0). "
    "S1/S2 baselines must be built so H2 (orchestration) is answerable. Everything else is ready.",
    size=12.5, color=RED, bold=True, spacing=1.15)
notes(s, "One-slide summary of the audit and plan.")

# ══════════════════════════════════════════════════════════════════════════
# 03 AGENDA
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("Agenda", "Overview")
items = [
    ("1", "Audit — what exists", "System architecture, code scale, benchmark inventory, capability readiness"),
    ("2", "Real execution evidence", "closed48 run, traces, latency, evidence depth — and what they do/don't prove"),
    ("3", "The experiment", "H1–H6, systems S0–S3, matrix, metrics, grading, statistics"),
    ("4", "End-to-end & failure safety", "G0–G12 workflow, fault injection, causal ablation, reproducibility"),
    ("5", "Feasibility & rigor", "cost–benefit, risks, go/no-go, falsification, roadmap"),
]
y = Inches(1.65)
for num, head, sub in items:
    _rect(s, Inches(0.6), y, Inches(12.1), Inches(1.0), PAPER, line=LINE, radius=0.07)
    _rect(s, Inches(0.6), y, Inches(0.95), Inches(1.0), NAVY, radius=0.07)
    _tb(s, Inches(0.6), y, Inches(0.95), Inches(1.0), num, size=24, color=WHITE,
        bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    _tb(s, Inches(1.75), y + Inches(0.13), Inches(10.7), Inches(0.4), head,
        size=16, color=NAVY, bold=True)
    _tb(s, Inches(1.75), y + Inches(0.55), Inches(10.7), Inches(0.4), sub,
        size=11.5, color=SLATE)
    y += Inches(1.08)
notes(s, "Roadmap.")

# ══════════════════════════════════════════════════════════════════════════
# 04 THE QUESTION
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("The question, and why it must be decomposed", "Framing")
_rect(s, Inches(0.6), Inches(1.45), Inches(12.1), Inches(0.95), NAVY, radius=0.08)
_tb(s, Inches(0.9), Inches(1.55), Inches(11.5), Inches(0.8),
    "H:  A structured TPD-specific scientific agent produces more scientifically correct, evidence-supported "
    "and executable PROTAC-design decisions than the underlying LLM alone.",
    size=14.5, color=WHITE, spacing=1.15)
bullets(s, [
    "H is too broad to test directly: it bundles scaffolding, retrieval, tools, workflow and evidence together.",
    "We decompose it into six experimentally separable hypotheses H1–H6.",
    "The base LLM is held constant — each model is compared to itself with vs without the scaffold.",
    "The experimental unit is the question/case, never the individual stochastic run.",
    "Primary endpoint and interpretation rules are frozen before any run (no metric shopping).",
], y=Inches(2.75), size=15, gap=12, h=Inches(4.0))
notes(s, "Decomposition principle.")

# ══════════════════════════════════════════════════════════════════════════
# 05 SECTION — AUDIT
# ══════════════════════════════════════════════════════════════════════════
section_slide("Part 1", "What we audited", "The real implementation, the real benchmark, the real blockers")

# ══════════════════════════════════════════════════════════════════════════
# 06 ARCHITECTURE
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "ProtacXtend as implemented", "System audit",
    "f06_architecture",
    "Request → planner → 34-node capability-routed graph → 15 specialist agents / ~30 callable tools → "
    "retrieval + 19 scientific backends → evidence ledger & state → ranking, gates, report.",
    "Not a scaffold: every box maps to real code (parser, graph.py, agents/*, agentic/registry.py, scientific_backends/).",
    layout="full", source="REAL — counts from source audit", source_color=TEAL)

# ══════════════════════════════════════════════════════════════════════════
# 07 MODULE SCALE
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "A real codebase, not a prompt wrapper", "System audit",
    "f07_module_scale",
    "Lines of Python by subsystem. The scientific reasoning machinery (agents, tools, modules, "
    "backends) is substantial — which is exactly why the orchestration effect (H2) is worth testing.",
    [("Tool registry: ~30 LLM-callable exec_* over 100 tool modules.", 0),
     ("27 TPD capabilities mapped to 19 backends with explicit fallback classes.", 0),
     ("8-scenario retrieval fault injector already implemented.", 0),
     ("Statistics, failure and reproducibility modules exist in tpdeval/.", 0)],
    layout="side", side="left", source="REAL — source tree", source_color=TEAL)

# ══════════════════════════════════════════════════════════════════════════
# 08 METRIC HIERARCHY
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "One primary endpoint, five secondaries, no collapsing", "Measurement design",
    "f18_metric_hierarchy",
    "A single composite can hide regressions. Scientific success is primary; evidence grounding, partial credit, "
    "evidence quality, abstention and tool use are reported separately, then end-to-end and safety endpoints.",
    layout="full", source="DESIGN", source_color=AMBER)

# ══════════════════════════════════════════════════════════════════════════
# 09 BENCHMARK INVENTORY
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "The benchmark exists — the gold does not", "Benchmark audit",
    "f01_benchmark_inventory",
    "Two 500-task banks and a 500-slot design manifest carry zero curated gold. The 48-case governed suite is the "
    "only asset with authored ground truth, yet 0/48 are independently adjudicated.",
    [("Correctness endpoints have denominator 0 until adjudication completes.", 0),
     ("19/48 governed cases are objectively scorable (5 exact, 10 categorical, 4 ranked).", 0),
     ("sota/eval is template material — smoke only, never a primary endpoint.", 0)],
    layout="side", side="right", source="REAL — files on disk", source_color=TEAL)

# ══════════════════════════════════════════════════════════════════════════
# 10 GOVERNED COMPOSITION
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "The 48 governed cases: where difficulty sits", "Benchmark audit",
    "f02_governed_composition",
    "Left: 12 cases each of KNOW / REASON / DESIGN / DISCOVER, concentrated at medium and hard. "
    "Right: ground-truth types — 19 are deterministic (teal), 29 are expert rubrics (amber).",
    "This is why grading must escalate from deterministic checks to blinded human adjudication.",
    layout="full", source="REAL — benchmark/ground_truth/*.json", source_color=TEAL,
    cap_size=12.5)

# ══════════════════════════════════════════════════════════════════════════
# 11 B500 DOMAINS
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "16 TPD domains, ~500 tasks — a rich but unscored bank", "Benchmark audit",
    "f03_b500_domains",
    "benchmark500 general_500 spans target biology, TPD tractability, E3 selection, warhead discovery, "
    "linker design, ternary reasoning, degradation, ADME, safety, resistance and translation.",
    [("Ideal source for a stratified POP-A once a subset is curated to gold.", 0),
     ("Identical domain taxonomy in the temporal suite enables POP-C.", 0),
     ("integrity-checked: 500 unique IDs, no missing fields, no duplicates.", 0)],
    layout="side", side="left", source="REAL — benchmark500", source_color=TEAL)

# ══════════════════════════════════════════════════════════════════════════
# 12 B500 DIFFICULTY
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "Difficulty is skewed toward reasoning, not recall", "Benchmark audit",
    "f04_b500_difficulty",
    "L1 retrieval → L7 discovery. Both suites peak at L4–L6, so a headline number will be dominated by "
    "mechanistic and design reasoning rather than lookup.",
    [("Report per-difficulty curves (supplement S1) to detect where gains vanish.", 0),
     ("Class imbalance across L1–L7 is a stated risk; macro averaging used.", 0)],
    layout="side", side="right", source="REAL — benchmark500", source_color=TEAL)

# ══════════════════════════════════════════════════════════════════════════
# 13 CAPABILITY READINESS
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "11 of 27 capabilities are blocked — and must not be faked", "System audit",
    "f05_capability_readiness",
    "Linker generation, de-novo generation, reaction prediction, protein language models, patent mining, "
    "proteomics, imaging and quantum chemistry are BLOCKED. The system must fail closed, never silently "
    "substitute a surrogate.",
    layout="full", source="REAL — config/capability_backend_crosswalk.yaml", source_color=TEAL)

# ══════════════════════════════════════════════════════════════════════════
# 14 SECTION — REAL EVIDENCE
# ══════════════════════════════════════════════════════════════════════════
section_slide("Part 2", "Real execution evidence", "What the system already does — and the limits of it")

# ══════════════════════════════════════════════════════════════════════════
# 15 TRACE EVENTS
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "The system leaves an auditable trace", "Real run data",
    "f11_trace_events",
    "Across 458 run traces: run lifecycle, 534 tool calls, node start/end timing, decisions and 28 errors. "
    "This is the raw material for tool-utilization and stage-path reproducibility metrics.",
    [("Tool selection, execution and utilization are all observable.", 0),
     ("Error events exist and are typed — needed for the failure experiment.", 0),
     ("No dependence on self-reported summaries.", 0)],
    layout="side", side="left", source="REAL — 458 outputs/runs traces", source_color=TEAL)

# ══════════════════════════════════════════════════════════════════════════
# 16 CLOSED48 OUTCOMES
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "A real 48-case × 3-arm run already exists", "Real run data",
    "f08_closed48_outcomes",
    "ProtacXtend completed 26/48, partially completed 6/48 and appropriately abstained on 16/48. "
    "The direct-tool arm completed 45/48; the fixed workflow produced 48 partials.",
    [("This is pipeline-integrity behaviour, NOT scientific correctness (gold pending).", 0),
     ("Typed abstention is already visible — directly relevant to H4.", 0),
     ("Becomes H1–H3 evidence only after adjudication binds the labels.", 0)],
    layout="side", side="right", source="REAL — closed48 retrieval_reliability run (n=144)", source_color=TEAL)

# ══════════════════════════════════════════════════════════════════════════
# 17 CLOSED48 LATENCY
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "Latency: the orchestration cost is real", "Real run data",
    "f09_closed48_latency",
    "Per-case latency distribution by arm (n=48 each). In this run the direct-tool arm is fastest; "
    "the full stack and the fixed workflow pay for coordination and retrieval.",
    [("Feeds the benefit/cost panel (H) and the cost–benefit interpretation rule.", 0),
     ("Runtime is an explicit secondary cost, not hidden.", 0)],
    layout="side", side="left", source="REAL — closed48 run", source_color=TEAL)

# ══════════════════════════════════════════════════════════════════════════
# 18 CLOSED48 EVIDENCE
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "Evidence citation depth differs by arm", "Real run data",
    "f10_closed48_evidence",
    "Mean evidence references per case (± SEM). Evidence grounding (H3) is measurable from the logs today — "
    "what is missing is the correctness label, not the evidence signal.",
    [("Supports building the claim-level evidence pipeline before gold lands.", 0),
     ("Combined with support labels, yields evidence-grounded success.", 0)],
    layout="side", side="right", source="REAL — closed48 run", source_color=TEAL)

# ══════════════════════════════════════════════════════════════════════════
# 19 NODE WATERFALL
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "Per-node timing inside a real workflow", "Real run data",
    "f12_node_waterfall",
    "Node elapsed times captured in an execution trace. This is the basis for the first-failure-stage and "
    "stage-path consistency measurements in the end-to-end experiment.",
    layout="full", source="REAL — outputs/runs/trace_test", source_color=TEAL)

# ══════════════════════════════════════════════════════════════════════════
# 20 INTERPRETATION / LIMITATION
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("What the real run proves — and what it does not", "Rigour")
_rect(s, Inches(0.5), Inches(1.45), Inches(6.05), Inches(4.9), TEAL_L, line=TEAL, radius=0.05)
_tb(s, Inches(0.75), Inches(1.6), Inches(5.6), Inches(0.5), "OBSERVATION (permitted today)",
    size=13, color=TEAL, bold=True)
bullets(s, [
    "The stack executes and produces typed, traceable outcomes on all 48 cases.",
    "ProtacXtend abstains on 16/48; direct tool on 2/48.",
    "Latency and evidence depth differ measurably by arm.",
    "Errors and retries are captured in the trace.",
], x=Inches(0.75), y=Inches(2.15), w=Inches(5.6), h=Inches(4.0), size=12.5, gap=10)
_rect(s, Inches(6.8), Inches(1.45), Inches(6.05), Inches(4.9), RED_L, line=RED, radius=0.05)
_tb(s, Inches(7.05), Inches(1.6), Inches(5.6), Inches(0.5), "NOT YET ESTABLISHED (needs gold)",
    size=13, color=RED, bold=True)
bullets(s, [
    "That any outcome is scientifically correct.",
    "That abstention was appropriate rather than a capability gap.",
    "That the orchestration architecture causes the difference (no matched S2).",
    "That the result is reproducible — this was a single pass, not replicates.",
], x=Inches(7.05), y=Inches(2.15), w=Inches(5.6), h=Inches(4.0), size=12.5, gap=10)
notes(s, "The honest reading: behavioural, not scientific, evidence.")

# ══════════════════════════════════════════════════════════════════════════
# 21 SECTION — THE EXPERIMENT
# ══════════════════════════════════════════════════════════════════════════
section_slide("Part 3", "The experiment", "Hypotheses, systems, populations, metrics, inference")

# ══════════════════════════════════════════════════════════════════════════
# 22 HYPOTHESES
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("Six hypotheses, each experimentally separable", "H1–H6")
table(s,
      ["ID", "Claim", "Isolates", "Contrast", "Falsified if"],
      [
          ["H1 Scaffold", "Agent beats plain LLM on correct task completion",
           "prompt + retrieval + tools + workflow vs none", "S3 vs S0", "Δ ≤ 0 / CI spans 0"],
          ["H2 Orchestration", "Workflow+agents beat flat tool access",
           "the architecture itself, tools held fixed", "S3 vs S2", "S3 ≈ S2"],
          ["H3 Evidence", "More supported claims, fewer unsupported / false citations",
           "grounding behaviour", "S3 vs S0/S2", "support precision not ↑"],
          ["H4 Reliability", "Detect → recover → fallback → abstain, not hallucinate",
           "failure handling", "fault episodes", "hallucination not ↓"],
          ["H5 End-to-end", "QA gains transfer to full target→degrader workflows",
           "workflow-level science", "S3 vs S0/S2 e2e", "valid completion not ↑"],
          ["H6 Reproducibility", "Replicates converge on evidence, ranking, decision",
           "stochastic stability", "within-system", "agreement ≈ chance"],
      ],
      y=Inches(1.4), h=Inches(4.4), col_widths=[1.6, 4.0, 3.4, 1.6, 1.9], fsize=10.5)
_tb(s, Inches(0.5), Inches(6.0), Inches(12.3), Inches(0.9),
    "Holding the base LLM constant throughout. Each hypothesis maps to ≥1 experiment and each experiment "
    "maps to ≥1 hypothesis (no orphans, no metric shopping).",
    size=12, color=SLATE, italic=True, spacing=1.15)
notes(s, "H1-H6 with the mechanism each isolates.")

# ══════════════════════════════════════════════════════════════════════════
# 23 SYSTEMS
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("Four systems, one model, matched tools", "S0–S3")
hdrs = ["ID", "System", "Adds over previous", "Code entry point", "Gap"]
rows = [
    ["S0", "LLM only", "— (question only)", "baselines.LLMBaseline", "none"],
    ["S1", "LLM + retrieval", "controlled literature / DB context", "baselines.RetrievalOnly + synthesis", "build LLM-RAG"],
    ["S2", "LLM + flat tools", "exactly S3's tool schemas; no graph/agents/state/gating",
     "adapters.tool_only + LLM loop", "build LLM-tools"],
    ["S3", "Full ProtacXtend", "planner + 15 agents + workflow + backends + evidence + fallback",
     "agents.runtime.run_protacpilot", "exists"],
]
table(s, hdrs, rows, y=Inches(1.4), h=Inches(2.6),
      col_widths=[0.8, 2.3, 3.6, 3.4, 1.6], fsize=11)
_rect(s, Inches(0.5), Inches(4.25), Inches(12.33), Inches(2.1), PAPER, line=LINE, radius=0.05)
_tb(s, Inches(0.75), Inches(4.4), Inches(11.8), Inches(0.4),
    "Mandatory fairness manifest — every axis on which S0–S3 could differ is recorded; equality is never assumed:",
    size=12, color=NAVY, bold=True)
_tb(s, Inches(0.75), Inches(4.9), Inches(11.9), Inches(1.3),
    "model · provider · context window · web access · tool access · database access · retry budget · compute · "
    "time allowance · token budget · temperature · top-p · seed · prompt version · tool-registry version · "
    "retrieval-corpus version · memory mode",
    size=12, color=INK, spacing=1.3)
_tb(s, Inches(0.75), Inches(6.0), Inches(11.9), Inches(0.4),
    "S2 is the pivotal baseline: same tools, no orchestration. If S3 ≈ S2, the architecture adds nothing.",
    size=11.5, color=RED, italic=True)
notes(s, "S2 is the crucial control. It must be built before Phase 2.")

# ══════════════════════════════════════════════════════════════════════════
# 24 EXPERIMENT MATRIX
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "Scale of the study", "Experimental matrix",
    "f13_experiment_matrix",
    "Runs by experiment. Replicates: 3 for task-level, 5 for end-to-end. Temperature 0, top-p 1.0, "
    "max-tokens fixed, seeds where supported, memory reset for independence.",
    [("Primary: 300 questions × 4 systems × 3 replicates.", 0),
     ("End-to-end: 48 cases × 4 systems × 5 replicates.", 0),
     ("Ablation and cross-model reuse a fixed 150-task subset.", 0)],
    layout="side", side="left", source="DESIGN", source_color=AMBER)

# ══════════════════════════════════════════════════════════════════════════
# 25 TRACEABILITY
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "Hypothesis × experiment traceability", "Experimental control",
    "f15_traceability",
    "Every hypothesis is exercised by at least one experiment and every experiment supports at least one "
    "hypothesis. This is the anti-metric-shopping guarantee.",
    layout="full", source="DESIGN", source_color=AMBER)

# ══════════════════════════════════════════════════════════════════════════
# 26 METRICS WITH DEFINITIONS
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("Metrics — exact definitions, fixed in advance", "Measurement")
rows = [
    ["Primary", "Macro Scientific Success", "(1/|F|) Σ_f mean_{q∈f,r} success(q,s,r)  ·  families equal weight"],
    ["Secondary 1", "Evidence-grounded success", "success × 1[all material claims supported]"],
    ["Secondary 2", "Partial score", "rubric 0.00 / 0.25 / 0.50 / 0.75 / 1.00"],
    ["Secondary 3", "Support precision", "supported / claims-carrying-evidence;  + unsupported, contradicted, false-citation rates"],
    ["Secondary 4", "Abstention quality", "appropriate abstention rate · false abstention rate"],
    ["Secondary 5", "Tool performance", "selection precision/recall · execution success · useful-result · redundant · incorrect"],
]
table(s, ["Tier", "Metric", "Definition"], rows, y=Inches(1.4), h=Inches(3.4),
      col_widths=[1.6, 3.2, 7.5], fsize=11)
_rect(s, Inches(0.5), Inches(5.0), Inches(6.05), Inches(1.6), PAPER, line=LINE, radius=0.05)
_tb(s, Inches(0.7), Inches(5.1), Inches(5.7), Inches(1.4),
    "End-to-end (POP-B)\nstage success · completion · valid completion (no critical upstream error) · "
    "first-failure stage · error-propagation · recovery · critical-error rate",
    size=11, color=INK, spacing=1.2)
_rect(s, Inches(6.8), Inches(5.0), Inches(6.03), Inches(1.6), PAPER, line=LINE, radius=0.05)
_tb(s, Inches(7.0), Inches(5.1), Inches(5.7), Inches(1.4),
    "Reproducibility (H6)\nanswer agreement · evidence & candidate Jaccard · Spearman/Kendall · "
    "stage-path consistency · mean ± SD · CV",
    size=11, color=INK, spacing=1.2)
notes(s, "Mathematical definitions. Correct-but-unsupported scores 0 on evidence-grounded success.")

# ══════════════════════════════════════════════════════════════════════════
# 27 GRADING
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("Grading — deterministic first, humans last", "Measurement")
ladder = [
    ("1", "Exact match", TEAL), ("2", "Synonym dict", TEAL), ("3", "Numeric tol.", TEAL),
    ("4", "Canonical SMILES", NAVY), ("5", "Database IDs", NAVY), ("6", "Ranking metrics", NAVY),
    ("7", "Rubric checklist", AMBER), ("8", "Blinded LLM judge", RED), ("9", "Human adjudication", RED),
]
x = Inches(0.5)
for num, lbl, col in ladder:
    _rect(s, x, Inches(1.5), Inches(1.34), Inches(1.55), col, radius=0.14)
    _tb(s, x, Inches(1.62), Inches(1.34), Inches(0.5), num, size=21, color=WHITE, bold=True,
        align=PP_ALIGN.CENTER)
    _tb(s, x + Inches(0.05), Inches(2.12), Inches(1.24), Inches(0.85), lbl, size=9.5,
        color=WHITE, bold=True, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    x += Inches(1.4)
bullets(s, [
    "Use the cheapest valid deterministic grader; escalate to an LLM judge only for free-text reasoning.",
    "Chemistry graded by RDKit canonical SMILES / InChIKey; identifiers by UniProt/PDB/ChEMBL normalisation.",
    "Blinding: judge and experts see randomised outputs with system / model / condition stripped.",
    "Two independent raters on a ~10% stratified sample; report Cohen's κ; adjudicate disagreements.",
    "Normative rule: a fabricated claim forces partial score 0.",
], y=Inches(3.4), size=14, gap=11, h=Inches(3.3))
notes(s, "Grading ladder and blinding protocol.")

# ══════════════════════════════════════════════════════════════════════════
# 28 STATISTICS
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("Statistical analysis plan", "Inference")
rows = [
    ["Primary S3 vs S0 (macro success)", "Paired bootstrap 95% CI, B = 10,000", "resample questions"],
    ["Paired binary outcomes", "McNemar exact test", "question level"],
    ["Paired partial scores", "Two-sided Wilcoxon signed-rank", "question level"],
    ["Task-family comparisons", "Per-family effect + Benjamini–Hochberg", "families"],
    ["Ablations", "Δ vs Full + 95% CI + Holm–Bonferroni", "ablations"],
    ["Sensitivity model", "success ~ system + difficulty + family + tool + depth + (1|question)", "mixed-effects logistic"],
    ["Robustness", "binary · partial · macro · micro · majority · mean replicate", "re-scoring"],
]
table(s, ["Comparison", "Test", "Unit / correction"], rows, y=Inches(1.45), h=Inches(4.1),
      col_widths=[4.7, 4.9, 2.7], fsize=11.5)
_tb(s, Inches(0.5), Inches(5.75), Inches(12.3), Inches(0.8),
    "Effect size + 95% CI + n always; never a p-value without an effect size. Replicates are averaged/majority-voted "
    "within a question — they are never independent samples.",
    size=12, color=SLATE, italic=True, spacing=1.2)
notes(s, "Paired, question-level, multiplicity-corrected.")

# ══════════════════════════════════════════════════════════════════════════
# 29 SECTION — E2E & SAFETY
# ══════════════════════════════════════════════════════════════════════════
section_slide("Part 4", "End-to-end & failure safety", "The two experiments that matter most")

# ══════════════════════════════════════════════════════════════════════════
# 30 WORKFLOW
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "The end-to-end PROTAC workflow (G0–G12)", "Hypothesis H5",
    "f16_workflow_g012",
    "Each case is scored stage-by-stage with codes PASS / PARTIAL / FAIL / UNSUPPORTED / NOT_REACHED / "
    "APPROPRIATE_ABSTENTION. This exposes error propagation visually — far more informative than aggregate accuracy.",
    "Critical metrics: stage success · completion · valid completion · first-failure stage · error-propagation rate · recovery rate · critical-error rate.",
    layout="full", source="DESIGN", source_color=AMBER)

# ══════════════════════════════════════════════════════════════════════════
# 31 FAILURE MATRIX
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "Does it fail safely?", "Hypothesis H4",
    "f17_failure_matrix",
    "15 controlled fault classes injected, each episode scored for detection, recovery, fallback, abstention and "
    "hallucinated continuation.",
    [("Hallucinated continuation is the headline safety endpoint.", 0),
     ("A system that invents continuity after a failure is unsafe, not merely wrong.", 0),
     ("Reuses the existing runtime/fault_injection.py harness.", 0)],
    layout="side", side="right", source="DESIGN", source_color=AMBER)

# ══════════════════════════════════════════════════════════════════════════
# 32 ABLATION
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("Causal ablation: why performance changes", "Hypothesis H2")
rows = [
    ["− RAG / retrieval", "memory/literature_rag.py, research/", "H3"],
    ["− specialist agents", "agents/*_agent.py (route subset)", "H2"],
    ["− workflow graph", "agents/graph.py CAPABILITY_NODES", "H2/H5"],
    ["− persistent state", "memory/run_state, WorkflowState", "H6"],
    ["− evidence gating", "evidence/*, design_gates.py", "H3/H4"],
    ["− scientific backends", "scientific_backends/ dispatch", "H1/H5"],
    ["− structural reasoning", "ternary_agent, structural/, docking", "H5"],
    ["− degradation prediction", "models/degradation_model.py", "H5"],
    ["− fallback / retry", "graph._should_retry, escalation/", "H4"],
    ["− TPD-specific planning", "CAPABILITY_NODES routing", "H2"],
]
table(s, ["Ablated component", "Code surface", "Hypothesis"], rows, y=Inches(1.4), h=Inches(4.5),
      col_widths=[3.6, 6.0, 1.7], fsize=11)
_tb(s, Inches(0.5), Inches(6.05), Inches(12.3), Inches(0.7),
    "Δ = Full − Ablation for scientific success, evidence-grounded success, e2e success, hallucination, runtime and "
    "tool calls — with effect size and uncertainty, never 'lower/higher' alone.",
    size=12, color=SLATE, italic=True, spacing=1.2)
notes(s, "Remove one component at a time; quantify effect size.")

# ══════════════════════════════════════════════════════════════════════════
# 33 REPRODUCIBILITY
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("Reproducibility: success or stochastic luck?", "Hypothesis H6")
cards = [
    ("Final-answer agreement", "share of replicates agreeing on the principal conclusion"),
    ("Evidence overlap", "mean pairwise Jaccard of evidence sets"),
    ("Candidate overlap", "mean pairwise Jaccard of top-k candidates"),
    ("Ranking reproducibility", "Spearman ρ and Kendall τ across replicates"),
    ("Stage-path consistency", "agreement in tool / agent execution sequence"),
    ("Score variability", "mean ± SD and coefficient of variation"),
]
x0, y0 = Inches(0.5), Inches(1.45)
for i, (h, b) in enumerate(cards):
    c, r = i % 3, i // 3
    x = x0 + c * Inches(4.14)
    y = y0 + r * Inches(1.95)
    _rect(s, x, y, Inches(3.94), Inches(1.75), PAPER, line=LINE, radius=0.06)
    _rect(s, x, y, Inches(3.94), Inches(0.1), TEAL, radius=0.4)
    _tb(s, x + Inches(0.25), y + Inches(0.28), Inches(3.45), Inches(0.6), h,
        size=14, color=NAVY, bold=True)
    _tb(s, x + Inches(0.25), y + Inches(0.88), Inches(3.45), Inches(0.8), b,
        size=11.5, color=INK, spacing=1.15)
_tb(s, Inches(0.5), Inches(5.55), Inches(12.3), Inches(0.9),
    "Independence guard: memory is reset between replicates; memory_mode is recorded per run. "
    "The frozen-state arm is used only in the persistent-state ablation. Report stability, not just the mean.",
    size=12, color=RED, italic=True, spacing=1.2)
notes(s, "Reproducibility endpoints and the memory-independence guard.")

# ══════════════════════════════════════════════════════════════════════════
# 34 COST
# ══════════════════════════════════════════════════════════════════════════
figure_slide(
    "Benefit / cost trade-off", "Feasibility",
    "f14_cost",
    "API cost is negligible relative to GPU wall-clock. Report the observed trade-off; do not declare "
    "unconditional superiority.",
    [("≈$147 API total vs ≈1,530 GPU-hours.", 0),
     ("End-to-end docking/MD dominates; decision point before Phase 3.", 0),
     ("Disk is 99% full: budget ≤60 GiB and prune raw e2e artifacts.", 0)],
    layout="side", side="left", source="ESTIMATE", source_color=AMBER)

# ══════════════════════════════════════════════════════════════════════════
# 35 RISKS
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("Experimental risks and controls", "Rigor")
rows = [
    ["Gold unadjudicated (0/48, 0/1000)", "benchmark audit", "Phase-0 gate; behavioural endpoints meanwhile"],
    ["Data leakage", "first pass clean, not exhaustive", "per-case tag; primary on clean subset"],
    ["Replicate dependence via memory", "persistent run_state is active", "memory reset; record per run"],
    ["S1/S2 not implemented", "no LLM+RAG, no LLM+tools loop", "build thin adapters; validate in pilot"],
    ["Judge bias", "LLM judge planned", "blind, randomise, 2 raters, report κ"],
    ["Backend instability (11/27 blocked)", "capability crosswalk", "fail closed; never silent surrogate"],
    ["Provider drift", "single provider (deepseek)", "pin model; re-smoke; record snapshot"],
    ["Disk exhaustion (99% full)", "165 GiB free", "prune raw artifacts; Parquet only"],
]
table(s, ["Risk", "Evidence", "Control"], rows, y=Inches(1.4), h=Inches(4.6),
      col_widths=[3.7, 3.7, 4.9], fsize=11)
notes(s, "Risks with controls; each ties to a gate.")

# ══════════════════════════════════════════════════════════════════════════
# 36 GO / NO-GO
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("Go / no-go criteria", "Decision gates")
_rect(s, Inches(0.5), Inches(1.45), Inches(6.05), Inches(4.9), TEAL_L, line=TEAL, radius=0.05)
_tb(s, Inches(0.75), Inches(1.6), Inches(5.6), Inches(0.5), "GO for pilot requires", size=16,
    color=NAVY, bold=True)
bullets(s, [
    "freeze manifest written (commit + tree hash)",
    "scoring plumbing passes offline smoke",
    "fault injector emits the 5 behavioural flags",
    "cost / latency / token logging works on one S3 run",
    "claim extractor runs on 20 answers",
    "GO for primary additionally: gold adjudicated (κ ≥ 0.7) and S1/S2 fairness-matched",
], x=Inches(0.75), y=Inches(2.25), w=Inches(5.6), h=Inches(3.9), size=12.5, gap=10)
_rect(s, Inches(6.8), Inches(1.45), Inches(6.03), Inches(4.9), RED_L, line=RED, radius=0.05)
_tb(s, Inches(7.05), Inches(1.6), Inches(5.6), Inches(0.5), "NO-GO triggers", size=16,
    color=RED, bold=True)
bullets(s, [
    "confirmed leakage in > 10% of POP-A",
    "S2 cannot be tool-matched to S3 → H2 unanswerable",
    "provider unavailable / model changed without re-smoke",
    "disk < 40 GiB free",
    "pilot shows scoring cannot discriminate systems",
], x=Inches(7.05), y=Inches(2.25), w=Inches(5.6), h=Inches(3.9), size=12.5, gap=10)
notes(s, "Gates that prevent an invalid study.")

# ══════════════════════════════════════════════════════════════════════════
# 37 FALSIFICATION
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("What would falsify the claim", "Scientific honesty")
items = [
    "Full ProtacXtend does not outperform the identical LLM with flat tool access (H2 fails).",
    "Gains disappear on difficult questions.",
    "Gains disappear on unseen targets.",
    "Valid end-to-end completion remains low.",
    "Agent orchestration increases hallucinated continuation (safety regression).",
    "Ablating the workflow architecture produces little or no change.",
    "Results are highly unstable across repeated runs.",
    "Gains arise primarily from benchmark leakage.",
]
y = Inches(1.5)
for i, it in enumerate(items):
    _rect(s, Inches(0.6), y, Inches(12.1), Inches(0.52), WHITE, line=LINE, radius=0.2)
    _tb(s, Inches(0.8), y, Inches(0.5), Inches(0.52), str(i + 1), size=13, color=RED,
        bold=True, anchor=MSO_ANCHOR.MIDDLE)
    _tb(s, Inches(1.3), y, Inches(11.2), Inches(0.52), it, size=12.5, color=INK,
        anchor=MSO_ANCHOR.MIDDLE)
    y += Inches(0.6)
_rect(s, Inches(0.6), Inches(6.35), Inches(12.1), Inches(0.62), NAVY, radius=0.08)
_tb(s, Inches(0.85), Inches(6.37), Inches(11.6), Inches(0.58),
    "These are first-class results and will be reported, not hidden.",
    size=13, color=WHITE, bold=True, anchor=MSO_ANCHOR.MIDDLE)
notes(s, "Pre-declared falsification conditions.")

# ══════════════════════════════════════════════════════════════════════════
# 38 ROADMAP
# ══════════════════════════════════════════════════════════════════════════
s = content_slide("Execution order", "Roadmap")
phases = [
    ("Phase 0", "Freeze · audits", "DONE", TEAL),
    ("Phase 1", "Pilot 30 × 4 × 3", "awaiting approval", AMBER),
    ("Phase 2", "Primary S0–S3", "gated on gold", SLATE),
    ("Phase 3", "End-to-end", "gated on gold", SLATE),
    ("Phase 4", "Failure injection", "ready", TEAL),
    ("Phase 5", "Ablations", "ready", TEAL),
    ("Phase 6", "Cross-model", "gated on providers", SLATE),
    ("Phase 7", "Statistics — freeze tables", "", SLATE),
    ("Phase 8", "Figures A–H + S1–S12", "", SLATE),
    ("Phase 9", "Interpretation & limits", "", SLATE),
]
x0, y0 = Inches(0.5), Inches(1.55)
for i, (ph, desc, state, col) in enumerate(phases):
    c, r = i % 2, i // 2
    x = x0 + c * Inches(6.25)
    y = y0 + r * Inches(1.02)
    _rect(s, x, y, Inches(6.0), Inches(0.86), WHITE, line=col, radius=0.08)
    _rect(s, x, y, Inches(0.12), Inches(0.86), col, radius=0.4)
    _tb(s, x + Inches(0.3), y, Inches(1.5), Inches(0.86), ph, size=13, color=NAVY,
        bold=True, anchor=MSO_ANCHOR.MIDDLE)
    _tb(s, x + Inches(1.8), y, Inches(2.9), Inches(0.86), desc, size=11.5, color=INK,
        anchor=MSO_ANCHOR.MIDDLE)
    _tb(s, x + Inches(4.65), y, Inches(1.25), Inches(0.86), state, size=9, color=col,
        bold=True, align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
_tb(s, Inches(0.5), Inches(6.75), Inches(12.3), Inches(0.4),
    "Result tables are frozen before any figure styling; figures are generated only from frozen tables.",
    size=11, color=SLATE, italic=True)
notes(s, "Phase 0 done; Phase 1 pilot is next.")

# ══════════════════════════════════════════════════════════════════════════
# 39 CLOSING
# ══════════════════════════════════════════════════════════════════════════
s = prs.slides.add_slide(BLANK)
_rect(s, 0, 0, W, H, NAVY, shape=MSO_SHAPE.RECTANGLE)
_rect(s, 0, 0, Inches(0.22), H, TEAL, shape=MSO_SHAPE.RECTANGLE)
_tb(s, Inches(0.9), Inches(1.2), Inches(11.6), Inches(1.0),
    "A study that can come back negative", size=33, color=WHITE, bold=True)
_tb(s, Inches(0.9), Inches(2.4), Inches(11.4), Inches(2.0),
    "H1–H6 are pre-declared, the unit of analysis is the question, grading escalates from deterministic to "
    "blinded human, and the falsification conditions are fixed in advance.\n\n"
    "If the orchestration architecture adds nothing over flat tools, this design will show it.  "
    "If it does add value — correctness, evidence quality and safe failure — the result will be hard to dismiss.",
    size=16, color=RGBColor(0xCF, 0xDD, 0xE5), spacing=1.3)
_rect(s, Inches(0.9), Inches(4.55), Inches(3.0), Inches(0.05), AMBER, shape=MSO_SHAPE.RECTANGLE)
_tb(s, Inches(0.9), Inches(4.85), Inches(11.4), Inches(0.9),
    "Next: approve the design → run the Phase 1 pilot → PILOT_REVIEW.md against the go/no-go checklist.",
    size=15, color=AMBER, bold=True)
_footer(s, "closing", dark=True)
notes(s, "Closing and ask.")

prs.save(OUT)
print(f"wrote {OUT}  ({len(prs.slides._sldIdLst)} slides)")
